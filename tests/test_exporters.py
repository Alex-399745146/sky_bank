# tests/test_exporters.py
"""Тесты экспорта Excel-отчётов."""

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from sky_bank.analytics import Overview
from sky_bank.exporters import export_analysis_report


def test_export_analysis_report_creates_formatted_workbook(tmp_path: Path) -> None:
    """Экспорт создаёт Excel-файл с тремя листами отчёта."""
    output_path = tmp_path / "analysis.xlsx"

    overview = Overview(
        currency="RUB",
        transaction_count=2,
        card_count=1,
        category_count=2,
        period_start=datetime(2021, 1, 3, 10, 0),
        period_end=datetime(2021, 1, 5, 15, 30),
        income=1_000.0,
        expenses=100.0,
        balance=900.0,
    )

    categories = pd.DataFrame(
        [
            {
                "category": "Продукты",
                "expenses": 100.0,
            }
        ]
    )

    weekdays = pd.DataFrame(
        [
            {
                "weekday_number": 0,
                "weekday": "Понедельник",
                "operation_count": 1,
                "expenses": 100.0,
                "average_expense": 100.0,
            }
        ]
    )

    result_path = export_analysis_report(
        output_path,
        overview,
        categories,
        weekdays,
    )

    assert result_path == output_path
    assert output_path.exists()

    workbook = load_workbook(output_path)

    assert workbook.sheetnames == [
        "Сводка",
        "Категории",
        "Дни недели",
    ]

    summary = workbook["Сводка"]
    categories_sheet = workbook["Категории"]
    weekdays_sheet = workbook["Дни недели"]

    assert summary["A1"].value == "Показатель"
    assert summary["B1"].value == "Значение"
    assert summary["A2"].value == "Валюта"
    assert summary["B2"].value == "RUB"
    assert summary.freeze_panes == "A2"

    assert categories_sheet["A1"].value == "Категория"
    assert categories_sheet["B1"].value == "Расходы"
    assert categories_sheet["A2"].value == "Продукты"
    assert categories_sheet["B2"].value == 100.0

    assert weekdays_sheet["A1"].value == "День недели"
    assert weekdays_sheet["B1"].value == "Операций"
    assert weekdays_sheet["C1"].value == "Расходы"
    assert weekdays_sheet["D1"].value == "Средний расход"
    assert weekdays_sheet["A2"].value == "Понедельник"
