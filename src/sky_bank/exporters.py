# src/sky_bank/exporters.py
"""Экспорт аналитических отчётов Sky Bank в Excel."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sky_bank.analytics import Overview

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")


def _format_worksheet(worksheet) -> None:
    """Оформляет заголовки, ширину столбцов и фильтры листа."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)


def _build_summary_dataframe(overview: Overview) -> pd.DataFrame:
    """Преобразует сводку операций в таблицу для Excel."""
    return pd.DataFrame(
        {
            "Показатель": [
                "Валюта",
                "Период начала",
                "Период окончания",
                "Количество операций",
                "Количество карт",
                "Количество категорий",
                "Доходы",
                "Расходы",
                "Баланс",
            ],
            "Значение": [
                overview.currency,
                overview.period_start.strftime("%d.%m.%Y"),
                overview.period_end.strftime("%d.%m.%Y"),
                overview.transaction_count,
                overview.card_count,
                overview.category_count,
                overview.income,
                overview.expenses,
                overview.balance,
            ],
        }
    )


def export_analysis_report(
    output_path: Path,
    overview: Overview,
    categories: pd.DataFrame,
    weekdays: pd.DataFrame,
) -> Path:
    """Создаёт Excel-отчёт из сводки, категорий и расходов по дням недели."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    summary = _build_summary_dataframe(overview)

    categories_report = categories.rename(
        columns={
            "category": "Категория",
            "expenses": "Расходы",
        }
    )

    weekdays_report = weekdays.rename(
        columns={
            "weekday": "День недели",
            "operation_count": "Операций",
            "expenses": "Расходы",
            "average_expense": "Средний расход",
        }
    )[
        [
            "День недели",
            "Операций",
            "Расходы",
            "Средний расход",
        ]
    ]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Сводка", index=False)
        categories_report.to_excel(writer, sheet_name="Категории", index=False)
        weekdays_report.to_excel(writer, sheet_name="Дни недели", index=False)

        workbook = writer.book

        for worksheet in workbook.worksheets:
            _format_worksheet(worksheet)

        summary_worksheet = workbook["Сводка"]

        for row_number in range(2, summary_worksheet.max_row + 1):
            metric = summary_worksheet.cell(row=row_number, column=1).value

            if metric in {"Доходы", "Расходы", "Баланс"}:
                summary_worksheet.cell(
                    row=row_number,
                    column=2,
                ).number_format = '#,##0.00'

        categories_worksheet = workbook["Категории"]
        weekdays_worksheet = workbook["Дни недели"]

        for row in categories_worksheet.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].number_format = '#,##0.00'

        for row in weekdays_worksheet.iter_rows(min_row=2, min_col=3, max_col=4):
            for cell in row:
                cell.number_format = '#,##0.00'

    return output_path